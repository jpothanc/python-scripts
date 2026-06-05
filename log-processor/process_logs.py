#!/usr/bin/env python3
"""Process timestamp-only log files and export matched metrics to Excel."""

from __future__ import annotations

import argparse
import re
import statistics
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.child import INVALID_TITLE_REGEX
from openpyxl.worksheet.worksheet import Worksheet

MAX_LOG_FILES = 4

TIME_PATTERN = re.compile(
    r"(?<!\d)"
    r"(\d{1,2}):(\d{2}):(\d{2})"
    r"(?:[.,](\d{1,6}))?"
    r"(?!\d)"
)


PLACEHOLDER_SPLIT = re.compile(r"(?i)x")


@dataclass(frozen=True)
class LogEntry:
    entry_datetime: datetime
    seconds: float


@dataclass(frozen=True)
class BusinessHours:
    start: time
    end: time


DEFAULT_BUSINESS_HOURS = BusinessHours(start=time(9, 0), end=time(16, 0))


@dataclass(frozen=True)
class ProcessorConfig:
    start_date: date | None
    line_pattern: str
    value_divisor: float
    min_seconds: float
    business_hours: BusinessHours
    output_file: Path
    log_files: list[Path]


@dataclass(frozen=True)
class SourceInsights:
    counts: dict[str, int]
    bh_counts: dict[str, int]
    off_counts: dict[str, int]
    total: int
    bh_total: int
    off_total: int
    avg_seconds: float
    median_seconds: float
    p95_seconds: float
    max_seconds: float
    pct_over_1min: float
    pct_over_5min: float
    first_seen: datetime | None
    last_seen: datetime | None


@dataclass(frozen=True)
class SummaryContext:
    start_date: date
    line_pattern: str
    min_seconds: float
    business_hours: BusinessHours
    output_file: Path
    log_files: list[Path]
    generated_at: datetime


def pattern_to_regex(pattern: str) -> re.Pattern[str]:
    if not PLACEHOLDER_SPLIT.search(pattern):
        raise ValueError("line_pattern must contain 'x' as the numeric placeholder.")

    parts = PLACEHOLDER_SPLIT.split(pattern)
    regex_body = r"(\d+(?:\.\d+)?)".join(re.escape(part) for part in parts)
    return re.compile(regex_body, re.IGNORECASE)


def parse_time_from_line(line: str) -> time | None:
    match = TIME_PATTERN.search(line)
    if not match:
        return None

    hour, minute, second, fraction = match.groups()
    microsecond = 0
    if fraction:
        microsecond = int(fraction.ljust(6, "0")[:6])

    return time(
        hour=int(hour),
        minute=int(minute),
        second=int(second),
        microsecond=microsecond,
    )


def assign_dates(entries: list[tuple[time, float]], start_date: date) -> list[LogEntry]:
    if not entries:
        return []

    current_date = start_date
    previous_time: time | None = None
    results: list[LogEntry] = []

    for entry_time, seconds in entries:
        if previous_time is not None and entry_time < previous_time:
            current_date += timedelta(days=1)

        results.append(
            LogEntry(
                entry_datetime=datetime.combine(current_date, entry_time),
                seconds=seconds,
            )
        )
        previous_time = entry_time

    return results


def process_log_file(
    log_path: Path,
    line_pattern: re.Pattern[str],
    start_date: date,
    value_divisor: float,
    min_seconds: float,
) -> list[LogEntry]:
    raw_entries: list[tuple[time, float]] = []

    with log_path.open(encoding="utf-8", errors="replace") as log_file:
        for line in log_file:
            pattern_match = line_pattern.search(line)
            if not pattern_match:
                continue

            entry_time = parse_time_from_line(line)
            if entry_time is None:
                continue

            raw_value = float(pattern_match.group(1))
            raw_entries.append((entry_time, raw_value / value_divisor))

    entries = assign_dates(raw_entries, start_date)
    if min_seconds <= 0:
        return entries
    return [entry for entry in entries if entry.seconds > min_seconds]


def sanitize_sheet_name(file_path: Path, used_names: set[str]) -> str:
    base_name = file_path.stem[:31]
    base_name = INVALID_TITLE_REGEX.sub("_", base_name).strip("'") or "Sheet"

    candidate = base_name
    suffix = 1
    while candidate in used_names:
        tail = f"_{suffix}"
        candidate = f"{base_name[: 31 - len(tail)]}{tail}"
        suffix += 1

    used_names.add(candidate)
    return candidate


ALL_SHEET_NAME = "All"
SUMMARY_SHEET_NAME = "Summary"

# Summary sheet styling — soft neutrals, readable in Excel
SUMMARY_TITLE_FILL = PatternFill("solid", fgColor="F1F5F9")
SUMMARY_TITLE_FONT = Font(bold=True, size=13, color="334155")
SUMMARY_LABEL_FILL = PatternFill("solid", fgColor="F8FAFC")
SUMMARY_VALUE_FILL = PatternFill("solid", fgColor="FFFFFF")
SUMMARY_LABEL_FONT = Font(bold=True, color="64748B", size=10)
SUMMARY_VALUE_FONT = Font(color="334155", size=10)
TABLE_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
TABLE_HEADER_FONT = Font(bold=True, color="475569", size=10)
TABLE_ROW_FILL_A = PatternFill("solid", fgColor="FFFFFF")
TABLE_ROW_FILL_B = PatternFill("solid", fgColor="F9FAFB")
TABLE_TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
TABLE_TOTAL_FONT = Font(bold=True, color="334155", size=10)
TABLE_DATA_FONT = Font(color="334155", size=10)
SECTION_TITLE_FILL = PatternFill("solid", fgColor="F1F5F9")
SECTION_TITLE_FONT = Font(bold=True, size=11, color="475569")
SOFT_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)

BUCKET_LABELS: tuple[str, ...] = (
    "<=10s",
    ">10s",
    ">30s",
    ">1min",
    ">2min",
    ">3min",
    ">4min",
    ">5min",
)


def duration_bucket(seconds: float) -> str:
    """Single mutually exclusive bucket label for Excel filtering."""
    if seconds > 300:
        return ">5min"
    if seconds > 240:
        return ">4min"
    if seconds > 180:
        return ">3min"
    if seconds > 120:
        return ">2min"
    if seconds > 60:
        return ">1min"
    if seconds > 30:
        return ">30s"
    if seconds > 10:
        return ">10s"
    return "<=10s"


def parse_clock_time(value: str, field_name: str) -> time:
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(str(value).strip(), fmt).time()
        except ValueError:
            continue
    raise ValueError(
        f"{field_name} must be a time like 09:00 or 16:00, got '{value}'."
    )


def parse_business_hours(settings: dict) -> BusinessHours:
    raw = settings.get("business_hours", {})
    if raw is None:
        return DEFAULT_BUSINESS_HOURS
    if not isinstance(raw, dict):
        raise ValueError("'settings.business_hours' must be a mapping.")

    start_raw = raw.get("start", "09:00")
    end_raw = raw.get("end", "16:00")
    return BusinessHours(
        start=parse_clock_time(str(start_raw), "business_hours.start"),
        end=parse_clock_time(str(end_raw), "business_hours.end"),
    )


def is_business_hours(entry_datetime: datetime, business_hours: BusinessHours) -> bool:
    entry_time = entry_datetime.time()
    start = business_hours.start
    end = business_hours.end
    if start <= end:
        return start <= entry_time <= end
    return entry_time >= start or entry_time <= end


def business_hours_flag(entry: LogEntry, business_hours: BusinessHours) -> str:
    return "Y" if is_business_hours(entry.entry_datetime, business_hours) else "N"


def format_business_hours_window(business_hours: BusinessHours) -> str:
    return (
        f"{business_hours.start.strftime('%H:%M')}"
        f" - {business_hours.end.strftime('%H:%M')}"
    )


def entry_row_values(entry: LogEntry, business_hours: BusinessHours) -> list:
    return [
        entry.entry_datetime,
        round(entry.seconds, 3),
        duration_bucket(entry.seconds),
        business_hours_flag(entry, business_hours),
    ]


def cell_display_length(value) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime):
        return 19
    if isinstance(value, float):
        return len(f"{value:.3f}")
    return len(str(value))


def fit_column_widths(
    sheet: Worksheet,
    *,
    first_col: int = 1,
    last_col: int | None = None,
    first_row: int = 1,
    last_row: int | None = None,
    min_widths: dict[int, float] | None = None,
    max_width: float = 52,
    padding: int = 2,
) -> None:
    last_col = last_col or sheet.max_column
    last_row = last_row or sheet.max_row
    min_widths = min_widths or {}

    for col in range(first_col, last_col + 1):
        longest = min_widths.get(col, 8)
        for row in range(first_row, last_row + 1):
            longest = max(longest, cell_display_length(sheet.cell(row=row, column=col).value))
        sheet.column_dimensions[get_column_letter(col)].width = min(
            max_width, longest + padding
        )


def summary_table_min_widths(headers: list[str]) -> dict[int, float]:
    widths: dict[int, float] = {}
    for col, header in enumerate(headers, start=1):
        if header == "Source":
            widths[col] = 14
        elif header in ("Details", "BH Details"):
            widths[col] = 30
        elif header in ("BH Y", "BH N"):
            widths[col] = 8
        elif header == "Business Hours":
            widths[col] = 14
        elif header in ("First seen", "Last seen"):
            widths[col] = 20
        elif header == "Total":
            widths[col] = 7
        elif header in BUCKET_LABELS:
            widths[col] = 9
        elif header.startswith("%"):
            widths[col] = 10
        elif header.endswith("(s)"):
            widths[col] = 11
        else:
            widths[col] = len(header) + 2
    return widths


def apply_datetime_format(sheet: Worksheet, datetime_column: int) -> None:
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=datetime_column).number_format = "yyyy-mm-dd hh:mm:ss"


def style_bucket_column(sheet: Worksheet, bucket_column: int) -> None:
    center = Alignment(horizontal="center")
    for row in range(1, sheet.max_row + 1):
        sheet.cell(row=row, column=bucket_column).alignment = center


def apply_sheet_layout(
    sheet: Worksheet,
    *,
    datetime_column: int,
    bucket_column: int,
    business_hours_column: int | None = None,
    datetime_width: int = 22,
    seconds_width: int = 12,
    bucket_width: int = 10,
) -> None:
    apply_datetime_format(sheet, datetime_column)
    style_bucket_column(sheet, bucket_column)
    if business_hours_column is not None:
        style_bucket_column(sheet, business_hours_column)
    sheet.column_dimensions[get_column_letter(datetime_column)].width = datetime_width
    sheet.column_dimensions[get_column_letter(bucket_column - 1)].width = seconds_width
    sheet.column_dimensions[get_column_letter(bucket_column)].width = bucket_width
    if business_hours_column is not None:
        sheet.column_dimensions[get_column_letter(business_hours_column)].width = 14
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def write_sheet(
    workbook: Workbook,
    sheet_name: str,
    entries: list[LogEntry],
    business_hours: BusinessHours,
) -> Worksheet:
    sheet = workbook.create_sheet(title=sheet_name)
    sheet.append(["DateTime", "Seconds", "Bucket", "Business Hours"])

    for entry in entries:
        sheet.append(entry_row_values(entry, business_hours))

    apply_sheet_layout(
        sheet,
        datetime_column=1,
        bucket_column=3,
        business_hours_column=4,
    )
    fit_column_widths(
        sheet,
        min_widths={1: 20, 2: 10, 3: 9, 4: 14},
    )
    return sheet


def count_buckets(entries: list[LogEntry]) -> dict[str, int]:
    counts = dict.fromkeys(BUCKET_LABELS, 0)
    for entry in entries:
        label = duration_bucket(entry.seconds)
        if label in counts:
            counts[label] += 1
    return counts


def format_bucket_details(counts: dict[str, int]) -> str:
    parts = [f"{label}: {counts[label]} times" for label in BUCKET_LABELS if counts[label] > 0]
    return ", ".join(parts) if parts else "no entries"


def count_buckets_for_entries(
    entries: list[LogEntry],
    *,
    business_hours: BusinessHours | None = None,
    within_business_hours: bool | None = None,
) -> dict[str, int]:
    counts = dict.fromkeys(BUCKET_LABELS, 0)
    for entry in entries:
        if within_business_hours is not None and business_hours is not None:
            in_bh = is_business_hours(entry.entry_datetime, business_hours)
            if in_bh != within_business_hours:
                continue
        label = duration_bucket(entry.seconds)
        if label in counts:
            counts[label] += 1
    return counts


def format_bucket_details_with_bh(
    all_counts: dict[str, int],
    bh_counts: dict[str, int],
) -> str:
    parts: list[str] = []
    for label in BUCKET_LABELS:
        total = all_counts[label]
        if total == 0:
            continue
        bh = bh_counts[label]
        off = total - bh
        parts.append(f"{label}: {total} (BH: {bh}, off: {off})")
    return ", ".join(parts) if parts else "no entries"


def percentile(seconds: list[float], pct: float) -> float:
    if not seconds:
        return 0.0
    if len(seconds) == 1:
        return seconds[0]
    ordered = sorted(seconds)
    rank = (len(ordered) - 1) * pct / 100
    low = int(rank)
    high = min(low + 1, len(ordered) - 1)
    weight = rank - low
    return ordered[low] * (1 - weight) + ordered[high] * weight


def build_source_insights(
    entries: list[LogEntry],
    business_hours: BusinessHours,
) -> SourceInsights:
    if not entries:
        empty_counts = dict.fromkeys(BUCKET_LABELS, 0)
        return SourceInsights(
            counts=empty_counts,
            bh_counts=empty_counts.copy(),
            off_counts=empty_counts.copy(),
            total=0,
            bh_total=0,
            off_total=0,
            avg_seconds=0.0,
            median_seconds=0.0,
            p95_seconds=0.0,
            max_seconds=0.0,
            pct_over_1min=0.0,
            pct_over_5min=0.0,
            first_seen=None,
            last_seen=None,
        )

    durations = [entry.seconds for entry in entries]
    datetimes = [entry.entry_datetime for entry in entries]
    total = len(entries)
    bh_total = sum(
        1 for entry in entries if is_business_hours(entry.entry_datetime, business_hours)
    )
    over_1min = sum(1 for value in durations if value > 60)
    over_5min = sum(1 for value in durations if value > 300)

    return SourceInsights(
        counts=count_buckets_for_entries(entries),
        bh_counts=count_buckets_for_entries(
            entries,
            business_hours=business_hours,
            within_business_hours=True,
        ),
        off_counts=count_buckets_for_entries(
            entries,
            business_hours=business_hours,
            within_business_hours=False,
        ),
        total=total,
        bh_total=bh_total,
        off_total=total - bh_total,
        avg_seconds=statistics.mean(durations),
        median_seconds=statistics.median(durations),
        p95_seconds=percentile(durations, 95),
        max_seconds=max(durations),
        pct_over_1min=round(100 * over_1min / total, 1),
        pct_over_5min=round(100 * over_5min / total, 1),
        first_seen=min(datetimes),
        last_seen=max(datetimes),
    )


def style_cell(
    cell,
    *,
    fill: PatternFill | None = None,
    font: Font | None = None,
    alignment: Alignment | None = None,
    border: Border | None = None,
) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border


def write_summary_info_block(sheet: Worksheet, context: SummaryContext) -> int:
    """Write run context at the top; return the row number where the data table starts."""
    sheet.merge_cells("A1:B1")
    title_cell = sheet["A1"]
    title_cell.value = "Log Analysis Summary"
    style_cell(
        title_cell,
        fill=SUMMARY_TITLE_FILL,
        font=SUMMARY_TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=SOFT_BORDER,
    )
    sheet.row_dimensions[1].height = 24

    min_filter = (
        f">{context.min_seconds}s only"
        if context.min_seconds > 0
        else "none (all matching lines)"
    )
    log_file_names = ", ".join(path.name for path in context.log_files)

    info_rows = [
        ("Generated", context.generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Start date", context.start_date.isoformat()),
        ("Business hours", format_business_hours_window(context.business_hours)),
        ("Line pattern", context.line_pattern),
        ("Min filter", min_filter),
        ("Output file", context.output_file.name),
        ("Log files", log_file_names),
    ]

    for index, (label, value) in enumerate(info_rows):
        row = index + 2
        label_cell = sheet.cell(row=row, column=1, value=label)
        value_cell = sheet.cell(row=row, column=2, value=value)
        style_cell(
            label_cell,
            fill=SUMMARY_LABEL_FILL,
            font=SUMMARY_LABEL_FONT,
            alignment=Alignment(horizontal="left", vertical="center", indent=1),
            border=SOFT_BORDER,
        )
        style_cell(
            value_cell,
            fill=SUMMARY_VALUE_FILL,
            font=SUMMARY_VALUE_FONT,
            alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
            border=SOFT_BORDER,
        )
        sheet.row_dimensions[row].height = 20

    info_end_row = len(info_rows) + 1
    fit_column_widths(
        sheet,
        first_col=1,
        last_col=2,
        first_row=1,
        last_row=info_end_row,
        min_widths={1: 14, 2: 24},
        max_width=60,
    )

    return len(info_rows) + 3


def write_all_sheet(
    workbook: Workbook,
    rows: list[tuple[str, LogEntry]],
    business_hours: BusinessHours,
) -> Worksheet:
    sheet = workbook.create_sheet(title=ALL_SHEET_NAME, index=0)
    sheet.append(["Source", "DateTime", "Seconds", "Bucket", "Business Hours"])

    for source, entry in sorted(rows, key=lambda item: (item[1].entry_datetime, item[0])):
        sheet.append([source, *entry_row_values(entry, business_hours)])

    apply_sheet_layout(
        sheet,
        datetime_column=2,
        bucket_column=4,
        business_hours_column=5,
        datetime_width=20,
        seconds_width=10,
        bucket_width=9,
    )
    fit_column_widths(
        sheet,
        min_widths={1: 14, 2: 20, 3: 10, 4: 9, 5: 14},
    )

    return sheet


def write_section_title(
    sheet: Worksheet,
    row: int,
    title: str,
    *,
    column_span: int = 10,
) -> None:
    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=column_span,
    )
    cell = sheet.cell(row=row, column=1, value=title)
    style_cell(
        cell,
        fill=SECTION_TITLE_FILL,
        font=SECTION_TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=SOFT_BORDER,
    )
    sheet.row_dimensions[row].height = 22


def write_table_header_row(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=header)
        style_cell(
            cell,
            fill=TABLE_HEADER_FILL,
            font=TABLE_HEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=SOFT_BORDER,
        )
    sheet.row_dimensions[row].height = 22


def write_styled_data_row(
    sheet: Worksheet,
    row: int,
    values: list,
    headers: list[str],
    *,
    is_total: bool = False,
    stripe_index: int = 0,
) -> None:
    center = Alignment(horizontal="center", vertical="center")
    row_fill = TABLE_TOTAL_FILL if is_total else (
        TABLE_ROW_FILL_B if stripe_index % 2 else TABLE_ROW_FILL_A
    )
    row_font = TABLE_TOTAL_FONT if is_total else TABLE_DATA_FONT

    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        header = headers[col - 1] if col - 1 < len(headers) else ""
        alignment = (
            Alignment(horizontal="left", vertical="center", indent=1)
            if col == 1
            else center
        )
        style_cell(
            cell,
            fill=row_fill,
            font=row_font,
            alignment=alignment,
            border=SOFT_BORDER,
        )
        if header in ("First seen", "Last seen") and value is not None:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.row_dimensions[row].height = 20


def write_performance_section(
    sheet: Worksheet,
    start_row: int,
    source_entries: list[tuple[str, list[LogEntry]]],
    context: SummaryContext,
) -> int:
    headers = [
        "Source",
        "Total",
        "BH Y",
        "BH N",
        "Avg (s)",
        "Median (s)",
        "P95 (s)",
        "Max (s)",
        "%>1 min",
        "%>5 min",
        "First seen",
        "Last seen",
    ]
    write_section_title(sheet, start_row, "Performance summary", column_span=len(headers))
    header_row = start_row + 1
    write_table_header_row(sheet, header_row, headers)

    data_row = header_row + 1
    for index, (source, entries) in enumerate(source_entries):
        insights = build_source_insights(entries, context.business_hours)
        write_styled_data_row(
            sheet,
            data_row,
            [
                source,
                insights.total,
                insights.bh_total,
                insights.off_total,
                round(insights.avg_seconds, 3),
                round(insights.median_seconds, 3),
                round(insights.p95_seconds, 3),
                round(insights.max_seconds, 3),
                insights.pct_over_1min,
                insights.pct_over_5min,
                insights.first_seen,
                insights.last_seen,
            ],
            headers,
            stripe_index=index,
        )
        data_row += 1

    if len(source_entries) > 1:
        all_entries = [entry for _, entries in source_entries for entry in entries]
        combined = build_source_insights(all_entries, context.business_hours)
        write_styled_data_row(
            sheet,
            data_row,
            [
                "ALL SOURCES",
                combined.total,
                combined.bh_total,
                combined.off_total,
                round(combined.avg_seconds, 3),
                round(combined.median_seconds, 3),
                round(combined.p95_seconds, 3),
                round(combined.max_seconds, 3),
                combined.pct_over_1min,
                combined.pct_over_5min,
                combined.first_seen,
                combined.last_seen,
            ],
            headers,
            is_total=True,
        )
        data_row += 1

    return data_row


def write_bucket_count_section(
    sheet: Worksheet,
    start_row: int,
    section_title: str,
    source_entries: list[tuple[str, list[LogEntry]]],
    context: SummaryContext,
    *,
    within_business_hours: bool,
) -> int:
    headers = ["Source", *BUCKET_LABELS, "Total"]
    write_section_title(sheet, start_row, section_title, column_span=len(headers))
    header_row = start_row + 1
    write_table_header_row(sheet, header_row, headers)

    data_row = header_row + 1
    for index, (source, entries) in enumerate(source_entries):
        insights = build_source_insights(entries, context.business_hours)
        counts = insights.bh_counts if within_business_hours else insights.off_counts
        total = insights.bh_total if within_business_hours else insights.off_total
        write_styled_data_row(
            sheet,
            data_row,
            [source, *[counts[label] for label in BUCKET_LABELS], total],
            headers,
            stripe_index=index,
        )
        data_row += 1

    if len(source_entries) > 1:
        all_entries = [entry for _, entries in source_entries for entry in entries]
        combined = build_source_insights(all_entries, context.business_hours)
        counts = combined.bh_counts if within_business_hours else combined.off_counts
        total = combined.bh_total if within_business_hours else combined.off_total
        write_styled_data_row(
            sheet,
            data_row,
            ["ALL SOURCES", *[counts[label] for label in BUCKET_LABELS], total],
            headers,
            is_total=True,
        )
        data_row += 1

    return data_row


def write_bucket_summary_sheet(
    workbook: Workbook,
    source_entries: list[tuple[str, list[LogEntry]]],
    context: SummaryContext,
) -> Worksheet:
    sheet = workbook.create_sheet(title=SUMMARY_SHEET_NAME, index=0)
    sheet.sheet_properties.tabColor = "94A3B8"
    next_row = write_summary_info_block(sheet, context)

    next_row = write_performance_section(sheet, next_row, source_entries, context)
    next_row += 1

    bh_window = format_business_hours_window(context.business_hours)
    next_row = write_bucket_count_section(
        sheet,
        next_row,
        f"Bucket counts — business hours ({bh_window})",
        source_entries,
        context,
        within_business_hours=True,
    )
    next_row += 1

    next_row = write_bucket_count_section(
        sheet,
        next_row,
        f"Bucket counts — outside business hours ({bh_window})",
        source_entries,
        context,
        within_business_hours=False,
    )

    fit_column_widths(sheet, min_widths={1: 14, 2: 9}, max_width=55)

    return sheet


def parse_start_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Invalid start date '{value}'. Use YYYY-MM-DD (e.g. 2026-06-01)."
        ) from exc


def resolve_path(path_value: str, base_dir: Path) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else base_dir / path


def load_config_file(config_path: Path) -> ProcessorConfig:
    with config_path.open(encoding="utf-8") as config_file:
        raw = yaml.safe_load(config_file)

    if not isinstance(raw, dict):
        raise ValueError("Config root must be a mapping.")

    settings = raw.get("settings", {})
    if not isinstance(settings, dict):
        raise ValueError("'settings' must be a mapping.")

    line_pattern = raw.get("line_pattern")
    if not line_pattern:
        raise ValueError("'line_pattern' is required in the config file.")

    log_files_raw = raw.get("log_files", [])
    if not isinstance(log_files_raw, list):
        raise ValueError("'log_files' must be a list when provided in config.")

    start_date_raw = settings.get("start_date")
    start_date = date.fromisoformat(str(start_date_raw)) if start_date_raw else None

    base_dir = config_path.parent
    output_raw = settings.get("output_file", "log_analysis.xlsx")
    output_file = resolve_path(str(output_raw), base_dir)

    value_divisor = float(raw.get("value_divisor", 1000))
    if value_divisor <= 0:
        raise ValueError("'value_divisor' must be greater than zero.")

    min_seconds = float(raw.get("min_seconds", 0))
    if min_seconds < 0:
        raise ValueError("'min_seconds' must be zero or greater.")

    log_files = [resolve_path(str(item), base_dir) for item in log_files_raw]

    business_hours = parse_business_hours(settings)

    return ProcessorConfig(
        start_date=start_date,
        line_pattern=str(line_pattern),
        value_divisor=value_divisor,
        min_seconds=min_seconds,
        business_hours=business_hours,
        output_file=output_file,
        log_files=log_files,
    )


def merge_config(
    file_config: ProcessorConfig | None,
    cli_start_date: date | None,
    cli_pattern: str | None,
    cli_output: str | None,
    cli_log_files: list[Path] | None,
    cli_min_seconds: float | None,
    config_base_dir: Path,
) -> ProcessorConfig:
    if file_config is None:
        if cli_start_date is None or not cli_pattern or not cli_log_files:
            raise ValueError(
                "Provide -c config.yaml, or pass --start-date, --pattern, and "
                "-f/--log-file (or positional log paths)."
            )
        output_file = Path(cli_output or "log_analysis.xlsx")
        if not output_file.is_absolute():
            output_file = config_base_dir / output_file
        return ProcessorConfig(
            start_date=cli_start_date,
            line_pattern=cli_pattern,
            value_divisor=1000.0,
            min_seconds=cli_min_seconds if cli_min_seconds is not None else 0.0,
            business_hours=DEFAULT_BUSINESS_HOURS,
            output_file=output_file,
            log_files=[path.resolve() for path in cli_log_files],
        )

    line_pattern = cli_pattern or file_config.line_pattern
    output_file = file_config.output_file
    if cli_output:
        output_file = Path(cli_output)
        if not output_file.is_absolute():
            output_file = config_base_dir / output_file

    log_files = (
        [path.resolve() for path in cli_log_files]
        if cli_log_files
        else file_config.log_files
    )
    if not log_files:
        raise ValueError(
            "At least one log file is required. Pass -f/--log-file or positional "
            "paths, or set log_files in the config file."
        )

    merged_start_date = cli_start_date or file_config.start_date
    if merged_start_date is None:
        raise ValueError(
            "Start date is required. Pass --start-date YYYY-MM-DD or set "
            "settings.start_date in the config file."
        )

    min_seconds = (
        cli_min_seconds if cli_min_seconds is not None else file_config.min_seconds
    )

    return ProcessorConfig(
        start_date=merged_start_date,
        line_pattern=line_pattern,
        value_divisor=file_config.value_divisor,
        min_seconds=min_seconds,
        business_hours=file_config.business_hours,
        output_file=output_file.resolve(),
        log_files=log_files,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Process up to four timestamp-only log files, extract a configurable "
            "line pattern, assign dates from a start date, and write one Excel workbook."
        )
    )
    parser.add_argument(
        "-c",
        "--config",
        type=Path,
        default=None,
        help="YAML config file (default: config.yaml in the current directory).",
    )
    parser.add_argument(
        "-d",
        "--start-date",
        type=parse_start_date,
        default=None,
        metavar="YYYY-MM-DD",
        help=(
            "First calendar date for log timestamps (e.g. 2026-06-01). "
            "Overrides settings.start_date in config when both are set."
        ),
    )
    parser.add_argument(
        "--pattern",
        default=None,
        help="Override config: line pattern with 'x' as the numeric placeholder.",
    )
    parser.add_argument(
        "--min-seconds",
        type=float,
        default=None,
        metavar="SECONDS",
        help=(
            "Only include lines where duration is greater than this many seconds. "
            "Overrides min_seconds in config (default: 0, include all)."
        ),
    )
    parser.add_argument(
        "-o",
        "--output",
        default=None,
        help="Override config: output Excel file path.",
    )
    parser.add_argument(
        "-f",
        "--log-file",
        action="append",
        type=Path,
        default=None,
        metavar="PATH",
        help=(
            "Log file to process (repeat up to 4 times). "
            "Overrides log_files in config when provided."
        ),
    )
    parser.add_argument(
        "log_files",
        nargs="*",
        type=Path,
        help="Log file paths (alternative to -f/--log-file). Up to four files.",
    )
    return parser.parse_args()


def collect_cli_log_files(args: argparse.Namespace) -> list[Path] | None:
    paths: list[Path] = []
    if args.log_file:
        paths.extend(args.log_file)
    if args.log_files:
        paths.extend(args.log_files)
    return paths or None


def main() -> int:
    args = parse_args()
    config_path = args.config

    if config_path is None:
        default_config = Path("config.yaml")
        if default_config.exists():
            config_path = default_config

    file_config: ProcessorConfig | None = None
    config_base_dir = Path.cwd()

    if config_path is not None:
        if not config_path.exists():
            print(f"Config file not found: {config_path}", file=sys.stderr)
            return 1
        try:
            file_config = load_config_file(config_path.resolve())
            config_base_dir = config_path.resolve().parent
        except (ValueError, yaml.YAMLError) as exc:
            print(f"Invalid config: {exc}", file=sys.stderr)
            return 1

    try:
        config = merge_config(
            file_config=file_config,
            cli_start_date=args.start_date,
            cli_pattern=args.pattern,
            cli_output=args.output,
            cli_log_files=collect_cli_log_files(args),
            cli_min_seconds=args.min_seconds,
            config_base_dir=config_base_dir,
        )
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 1

    if len(config.log_files) > MAX_LOG_FILES:
        print(
            f"At most {MAX_LOG_FILES} log files are supported; "
            f"received {len(config.log_files)}.",
            file=sys.stderr,
        )
        return 1

    try:
        line_pattern = pattern_to_regex(config.line_pattern)
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 1

    missing = [path for path in config.log_files if not path.exists()]
    if missing:
        for path in missing:
            print(f"Log file not found: {path}", file=sys.stderr)
        return 1

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_sheet_names: set[str] = {SUMMARY_SHEET_NAME, ALL_SHEET_NAME}
    total_rows = 0
    all_rows: list[tuple[str, LogEntry]] = []
    source_entries: list[tuple[str, list[LogEntry]]] = []

    for log_path in config.log_files:
        resolved = log_path.resolve()
        entries = process_log_file(
            resolved,
            line_pattern,
            config.start_date,
            config.value_divisor,
            config.min_seconds,
        )
        sheet_name = sanitize_sheet_name(resolved, used_sheet_names)
        write_sheet(workbook, sheet_name, entries, config.business_hours)
        source_label = resolved.stem
        source_entries.append((source_label, entries))
        all_rows.extend((source_label, entry) for entry in entries)
        total_rows += len(entries)
        print(f"{resolved.name}: {len(entries)} matching line(s) -> sheet '{sheet_name}'")

    if not all_rows:
        print("No matching lines found in any log file.", file=sys.stderr)
        return 1

    write_all_sheet(workbook, all_rows, config.business_hours)
    summary_context = SummaryContext(
        start_date=config.start_date,
        line_pattern=config.line_pattern,
        min_seconds=config.min_seconds,
        business_hours=config.business_hours,
        output_file=config.output_file,
        log_files=config.log_files,
        generated_at=datetime.now(),
    )
    write_bucket_summary_sheet(workbook, source_entries, summary_context)
    print(f"All: {len(all_rows)} row(s) -> sheet '{ALL_SHEET_NAME}'")
    print(f"Summary: bucket counts for {len(source_entries)} source(s) -> sheet '{SUMMARY_SHEET_NAME}'")

    workbook.save(config.output_file)
    min_filter = (
        f"min_seconds: >{config.min_seconds}"
        if config.min_seconds > 0
        else "min_seconds: none"
    )
    print(
        f"\nPattern: {config.line_pattern!r}"
        f"\n{min_filter}"
        f"\nWrote {config.output_file} ({len(workbook.sheetnames)} sheet(s), {total_rows} row(s))."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
