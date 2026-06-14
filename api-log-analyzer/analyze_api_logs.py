#!/usr/bin/env python3
"""Read API monitor CSV logs and export an Excel report with summary buckets."""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.child import INVALID_TITLE_REGEX
from openpyxl.worksheet.worksheet import Worksheet

RECORDS_SHEET = "Records"
SUMMARY_SHEET = "Summary"

# Mutually exclusive time buckets (milliseconds), plus timeout.
TIME_BUCKET_LABELS: tuple[str, ...] = (
    "<10ms",
    "10-20ms",
    "20-50ms",
    "50-100ms",
    "100-200ms",
    "200-500ms",
    "500ms-1s",
    "1s-1min",
    ">1min",
    "timeout",
)

TIMEOUT_FILL = PatternFill("solid", fgColor="FECACA")
TIMEOUT_FONT = Font(color="991B1B", bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
HEADER_FONT = Font(bold=True, color="475569")
SECTION_FILL = PatternFill("solid", fgColor="F1F5F9")
SECTION_FONT = Font(bold=True, size=11, color="475569")
ROW_FILL_A = PatternFill("solid", fgColor="FFFFFF")
ROW_FILL_B = PatternFill("solid", fgColor="F9FAFB")
TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
TOTAL_FONT = Font(bold=True, color="334155")
SOFT_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)


@dataclass(frozen=True)
class ApiRecord:
    timestamp: str
    api_name: str
    http_status: str
    time_ms: float
    result: str
    is_timeout: bool
    time_bucket: str


def parse_time_ms(value: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def is_timeout_row(result: str, http_status: str) -> bool:
    return result.strip().lower() == "timeout" or (
        not http_status.strip() and result.strip().lower() in {"timeout", "error"}
    )


def assign_time_bucket(time_ms: float, is_timeout: bool) -> str:
    if is_timeout:
        return "timeout"
    if time_ms < 10:
        return "<10ms"
    if time_ms <= 20:
        return "10-20ms"
    if time_ms <= 50:
        return "20-50ms"
    if time_ms <= 100:
        return "50-100ms"
    if time_ms <= 200:
        return "100-200ms"
    if time_ms <= 500:
        return "200-500ms"
    if time_ms <= 1000:
        return "500ms-1s"
    if time_ms <= 60000:
        return "1s-1min"
    return ">1min"


def load_records(log_path: Path) -> list[ApiRecord]:
    records: list[ApiRecord] = []

    with log_path.open(newline="", encoding="utf-8") as log_file:
        reader = csv.DictReader(log_file)
        if not reader.fieldnames:
            raise ValueError("Log file is empty or missing a header row.")

        required = {"api_name", "response_time_ms", "result"}
        missing = required - set(reader.fieldnames)
        if missing:
            raise ValueError(
                f"Log file missing required columns: {', '.join(sorted(missing))}"
            )

        for row in reader:
            api_name = (row.get("api_name") or "").strip()
            if not api_name:
                continue

            http_status = (row.get("http_status") or "").strip()
            result = (row.get("result") or "").strip()
            time_ms = parse_time_ms(row.get("response_time_ms", ""))
            timeout = is_timeout_row(result, http_status)

            records.append(
                ApiRecord(
                    timestamp=(row.get("timestamp_utc") or row.get("timestamp") or "").strip(),
                    api_name=api_name,
                    http_status=http_status if http_status else ("TIMEOUT" if timeout else ""),
                    time_ms=time_ms,
                    result=result or ("timeout" if timeout else "success"),
                    is_timeout=timeout,
                    time_bucket=assign_time_bucket(time_ms, timeout),
                )
            )

    return records


def empty_bucket_counts() -> dict[str, int]:
    return dict.fromkeys(TIME_BUCKET_LABELS, 0)


def count_buckets(records: list[ApiRecord]) -> dict[str, int]:
    counts = empty_bucket_counts()
    for record in records:
        counts[record.time_bucket] += 1
    return counts


def group_by_api(records: list[ApiRecord]) -> dict[str, list[ApiRecord]]:
    grouped: dict[str, list[ApiRecord]] = defaultdict(list)
    for record in records:
        grouped[record.api_name].append(record)
    return dict(sorted(grouped.items()))


def style_cell(
    cell,
    *,
    fill: PatternFill | None = None,
    font: Font | None = None,
    alignment: Alignment | None = None,
    border: Border | None = None,
) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border


def fit_column_widths(sheet: Worksheet, *, min_widths: dict[int, float] | None = None) -> None:
    min_widths = min_widths or {}
    for col in range(1, sheet.max_column + 1):
        longest = min_widths.get(col, 10)
        for row in range(1, sheet.max_row + 1):
            value = sheet.cell(row=row, column=col).value
            if value is not None:
                longest = max(longest, len(str(value)))
        sheet.column_dimensions[get_column_letter(col)].width = min(55, longest + 2)


def write_records_sheet(workbook: Workbook, records: list[ApiRecord]) -> None:
    sheet = workbook.create_sheet(title=RECORDS_SHEET)
    headers = [
        "Timestamp",
        "API Name",
        "HTTP Status",
        "Time (ms)",
        "Time Bucket",
        "Result",
    ]
    sheet.append(headers)

    center = Alignment(horizontal="center", vertical="center")
    for col, header in enumerate(headers, start=1):
        style_cell(
            sheet.cell(row=1, column=col, value=header),
            fill=HEADER_FILL,
            font=HEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center"),
            border=SOFT_BORDER,
        )

    for index, record in enumerate(records, start=2):
        row_values = [
            record.timestamp,
            record.api_name,
            record.http_status,
            round(record.time_ms, 2),
            record.time_bucket,
            record.result,
        ]
        row_fill = ROW_FILL_B if index % 2 else ROW_FILL_A
        for col, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=index, column=col, value=value)
            style_cell(
                cell,
                fill=TIMEOUT_FILL if record.is_timeout else row_fill,
                font=TIMEOUT_FONT if record.is_timeout else Font(color="334155"),
                alignment=center if col > 2 else Alignment(horizontal="left", vertical="center"),
                border=SOFT_BORDER,
            )

    sheet.auto_filter.ref = sheet.dimensions
    fit_column_widths(sheet, min_widths={1: 22, 2: 18, 3: 12, 4: 12, 5: 12, 6: 12})


def write_summary_sheet(workbook: Workbook, records: list[ApiRecord]) -> None:
    sheet = workbook.create_sheet(title=SUMMARY_SHEET, index=0)
    sheet.sheet_properties.tabColor = "94A3B8"

    sheet.merge_cells("A1:B1")
    style_cell(
        sheet["A1"],
        fill=SECTION_FILL,
        font=Font(bold=True, size=13, color="334155"),
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=SOFT_BORDER,
    )
    sheet["A1"] = "API response time summary"
    sheet.row_dimensions[1].height = 24

    info_rows = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total records", str(len(records))),
        ("Timeouts", str(sum(1 for record in records if record.is_timeout))),
    ]
    for index, (label, value) in enumerate(info_rows, start=2):
        style_cell(
            sheet.cell(row=index, column=1, value=label),
            fill=PatternFill("solid", fgColor="F8FAFC"),
            font=Font(bold=True, color="64748B"),
            border=SOFT_BORDER,
        )
        style_cell(
            sheet.cell(row=index, column=2, value=value),
            fill=PatternFill("solid", fgColor="FFFFFF"),
            font=Font(color="334155"),
            border=SOFT_BORDER,
        )

    table_start = len(info_rows) + 4
    headers = ["API Name", *TIME_BUCKET_LABELS, "Total"]
    for col, header in enumerate(headers, start=1):
        style_cell(
            sheet.cell(row=table_start, column=col, value=header),
            fill=HEADER_FILL,
            font=HEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=SOFT_BORDER,
        )

    center = Alignment(horizontal="center", vertical="center")
    data_row = table_start + 1
    grouped = group_by_api(records)

    for stripe, (api_name, api_records) in enumerate(grouped.items()):
        counts = count_buckets(api_records)
        row_values = [api_name, *[counts[label] for label in TIME_BUCKET_LABELS], len(api_records)]
        row_fill = ROW_FILL_B if stripe % 2 else ROW_FILL_A
        for col, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=data_row, column=col, value=value)
            highlight = col == headers.index("timeout") + 1 and value > 0
            style_cell(
                cell,
                fill=TIMEOUT_FILL if highlight else row_fill,
                font=TIMEOUT_FONT if highlight else Font(color="334155"),
                alignment=center if col > 1 else Alignment(horizontal="left", vertical="center"),
                border=SOFT_BORDER,
            )
        data_row += 1

    if len(grouped) > 1:
        total_counts = count_buckets(records)
        row_values = [
            "ALL APIS",
            *[total_counts[label] for label in TIME_BUCKET_LABELS],
            len(records),
        ]
        for col, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=data_row, column=col, value=value)
            highlight = col == headers.index("timeout") + 1 and value > 0
            style_cell(
                cell,
                fill=TIMEOUT_FILL if highlight else TOTAL_FILL,
                font=TIMEOUT_FONT if highlight else TOTAL_FONT,
                alignment=center if col > 1 else Alignment(horizontal="left", vertical="center"),
                border=SOFT_BORDER,
            )
        data_row += 1

    fit_column_widths(sheet, min_widths={1: 18})
    sheet.freeze_panes = sheet.cell(row=table_start + 1, column=1)


def load_paths_from_config(config_path: Path) -> tuple[Path, Path]:
    with config_path.open(encoding="utf-8") as config_file:
        raw = yaml.safe_load(config_file) or {}

    base_dir = config_path.parent
    input_file = Path(raw.get("input_file", "sample_logs/api_calls.csv"))
    output_file = Path(raw.get("output_file", "api_report.xlsx"))

    if not input_file.is_absolute():
        input_file = base_dir / input_file
    if not output_file.is_absolute():
        output_file = base_dir / output_file

    return input_file, output_file


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert API monitor CSV logs into an Excel report."
    )
    parser.add_argument(
        "-c",
        "--config",
        type=Path,
        default=None,
        help="YAML config file (default: config.yaml in the current directory).",
    )
    parser.add_argument(
        "-i",
        "--input",
        type=Path,
        default=None,
        help="Input CSV log file (overrides config).",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Output Excel file (overrides config).",
    )
    parser.add_argument(
        "log_file",
        nargs="?",
        type=Path,
        help="Optional input CSV path (alternative to -i).",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config_path = args.config
    if config_path is None and Path("config.yaml").exists():
        config_path = Path("config.yaml")

    input_path: Path | None = args.input or args.log_file
    output_path: Path | None = args.output

    if config_path is not None:
        if not config_path.exists():
            print(f"Config file not found: {config_path}", file=sys.stderr)
            return 1
        config_input, config_output = load_paths_from_config(config_path.resolve())
        input_path = input_path or config_input
        output_path = output_path or config_output

    if input_path is None:
        print("Provide -i/--input, a positional log file, or config.yaml with input_file.", file=sys.stderr)
        return 1

    output_path = (output_path or Path("api_report.xlsx")).resolve()
    input_path = input_path.resolve()

    if not input_path.exists():
        print(f"Log file not found: {input_path}", file=sys.stderr)
        return 1

    try:
        records = load_records(input_path)
    except (ValueError, csv.Error) as exc:
        print(f"Failed to read log file: {exc}", file=sys.stderr)
        return 1

    if not records:
        print("No records found in log file.", file=sys.stderr)
        return 1

    workbook = Workbook()
    workbook.remove(workbook.active)
    write_summary_sheet(workbook, records)
    write_records_sheet(workbook, records)
    workbook.save(output_path)

    timeouts = sum(1 for record in records if record.is_timeout)
    print(f"Read {len(records)} record(s) from {input_path}")
    print(f"Timeouts: {timeouts}")
    print(f"Wrote {output_path} (sheets: {SUMMARY_SHEET}, {RECORDS_SHEET})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
